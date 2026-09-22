"""
Script ULTIME : 
1. Fusionne les runs fragmentés
2. Convertit {xxx} en {{ xxx }} pour docxtpl/Jinja2
"""
import re
from docx import Document
from pathlib import Path


def fix_paragraph(paragraph, debug=True):
    """Fusionne les runs + convertit {xxx} en {{ xxx }}."""
    if not paragraph.runs:
        return 0
    
    # Concaténer tout le texte
    full_text = "".join(run.text for run in paragraph.runs)
    
    if '{' not in full_text and '}' not in full_text:
        return 0
    
    # ✅ Étape 1 : Convertir {xxx} en {{ xxx }}
    def to_double_braces(match):
        content = match.group(1).strip()
        return f"{{{{ {content} }}}}"
    
    new_text = re.sub(r'\{([^{}]+)\}', to_double_braces, full_text)
    
    if new_text == full_text and len(paragraph.runs) == 1:
        return 0
    
    # ✅ Étape 2 : Mettre tout dans le premier run et supprimer les autres
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)
    
    if debug:
        print(f"   ✅ {new_text[:100]}")
    
    return 1


def fix_table(table, debug=True):
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                count += fix_paragraph(paragraph, debug)
            for nested in cell.tables:
                count += fix_table(nested, debug)
    return count


def fix_template(template_path, debug=True):
    print(f"\n{'=' * 60}")
    print(f"🔧 Traitement : {template_path}")
    print(f"{'=' * 60}")
    
    doc = Document(str(template_path))
    total = 0
    
    # Corps
    for paragraph in doc.paragraphs:
        total += fix_paragraph(paragraph, debug)
    
    # Tableaux
    for table in doc.tables:
        total += fix_table(table, debug)
    
    # Headers/footers
    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header:
                for p in header.paragraphs:
                    total += fix_paragraph(p, debug)
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer:
                for p in footer.paragraphs:
                    total += fix_paragraph(p, debug)
    
    doc.save(str(template_path))
    print(f"\n✅ {total} paragraphe(s) réparé(s)\n")
    return total


def main():
    templates_dir = Path(__file__).parent / "templates"
    total = 0
    for name in ["convention_FR.docx", "convention_AR.docx"]:
        path = templates_dir / name
        if path.exists():
            total += fix_template(path, debug=True)
    print(f"\n🏁 TOTAL : {total} paragraphe(s)\n")


if __name__ == "__main__":
    main()